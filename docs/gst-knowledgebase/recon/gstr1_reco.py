#!/usr/bin/env python3
"""
GSTR-1 Books vs Portal reconciliation (Rules G1-G4, see SKILL_GSTR1_RECON.md).

Matches sales-register (books) invoices against the GST portal's GSTR-1 export
on: bill no, taxable value, GST rate, GST amount. Produces a working Excel with
SUMMARY / MISMATCH / BOOKS_ONLY / PORTAL_ONLY / MATCHED tabs.

Usage:
  python gstr1_reco.py books.xlsx portal.xlsx out_RECO.xlsx [--tol 1]
                       [--books-sheet NAME] [--portal-sheet NAME]

Column names are auto-detected (bill no / invoice number, GSTIN, taxable value,
rate, IGST/CGST/SGST or total tax). Header row is auto-found (portal exports
carry summary rows on top). Needs: pip install openpyxl
"""
import argparse
import csv
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

# ---------------------------------------------------------------- column map
SYNONYMS = {
    "inv":     ["invoice number", "invoice no", "invoice no.", "inv no", "bill no",
                "bill no.", "bill number", "document number", "doc no", "voucher no",
                "invoice num", "invoice"],
    "gstin":   ["gstin/uin of recipient", "gstin of recipient", "recipient gstin",
                "customer gstin", "party gstin", "buyer gstin", "gstin/uin", "gstin"],
    "date":    ["invoice date", "inv date", "bill date", "document date", "date"],
    "party":   ["receiver name", "party name", "customer name", "buyer name",
                "name of customer", "party"],
    "taxable": ["taxable value", "taxable amount", "taxable amt", "assessable value",
                "basic amount", "taxable"],
    "rate":    ["rate", "gst rate", "tax rate", "rate (%)", "gst %"],
    "igst":    ["igst amount", "integrated tax amount", "integrated tax", "igst"],
    "cgst":    ["cgst amount", "central tax amount", "central tax", "cgst"],
    "sgst":    ["sgst amount", "state/ut tax amount", "state tax amount", "state tax",
                "sgst/utgst", "sgst"],
    "totaltax": ["total tax amount", "total gst amount", "gst amount", "total tax",
                 "tax amount", "total gst"],
    "invval":  ["invoice value", "total invoice value", "bill amount", "gross total"],
}

def norm_header(h):
    return re.sub(r"\s+", " ", str(h)).strip().lower()

def map_columns(header):
    """header: list of cell values -> {field: col_index}. Exact synonym match,
    tried in SYNONYMS order (most specific first)."""
    hdrs = [norm_header(h) for h in header]
    out = {}
    for field, names in SYNONYMS.items():
        for name in names:
            if name in hdrs and hdrs.index(name) not in out.values():
                out[field] = hdrs.index(name)
                break
    return out

def find_header_row(rows, max_scan=15):
    """Portal exports have title/summary rows on top — find the real header."""
    best = (0, None, None)
    for i, row in enumerate(rows[:max_scan]):
        cols = map_columns(row)
        score = len(cols) + (2 if "inv" in cols else 0)
        if "inv" in cols and "taxable" in cols and score > best[0]:
            best = (score, i, cols)
    return best[1], best[2]

# ---------------------------------------------------------------- file load
def read_rows(path, sheet=None):
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig", errors="ignore") as f:
            return path.name, [list(r) for r in csv.reader(f)]
    wb = load_workbook(path, read_only=True, data_only=True)
    names = wb.sheetnames
    if sheet:
        pick = next((n for n in names if n.lower() == sheet.lower()), None)
        if not pick:
            sys.exit(f"Sheet '{sheet}' not in {path.name} (has: {', '.join(names)})")
        chosen = [pick]
    else:  # prefer a 'b2b' sheet (portal export), else scan all
        chosen = sorted(names, key=lambda n: (0 if "b2b" in n.lower() else 1))
    for name in chosen:
        rows = [[c for c in r] for r in wb[name].iter_rows(values_only=True)]
        if find_header_row(rows)[0] is not None:
            return f"{path.name}[{name}]", rows
    sys.exit(f"No sheet with invoice-no + taxable-value headers found in {path.name}")

def num(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^0-9.\-]", "", str(v))
    try:
        return float(s) if s not in ("", "-", ".") else 0.0
    except ValueError:
        return 0.0

# ---------------------------------------------------------------- G1: keys
def norm_inv(v):
    return re.sub(r"[^A-Z0-9]", "", str(v).upper())

def norm_inv_zeroless(v):
    # strip leading zeros inside numeric runs: INV0045 -> INV45
    return re.sub(r"\d+", lambda m: m.group().lstrip("0") or "0", norm_inv(v))

def load_side(label, source, rows):
    """Aggregate to invoice level. Returns {invkey: rec} keyed by normalized inv no."""
    hi, cols = find_header_row(rows)
    if hi is None:
        sys.exit(f"{label}: could not find header row in {source}")
    have_tax = any(f in cols for f in ("igst", "cgst", "sgst", "totaltax"))
    inv_map = {}
    for row in rows[hi + 1:]:
        if row is None or cols["inv"] >= len(row) or row[cols["inv"]] in (None, ""):
            continue
        raw_inv = str(row[cols["inv"]]).strip()
        if not norm_inv(raw_inv):
            continue
        g = lambda f: row[cols[f]] if f in cols and cols[f] < len(row) else None
        taxable = num(g("taxable"))
        rate = num(g("rate"))
        if have_tax:
            tax = num(g("totaltax")) or (num(g("igst")) + num(g("cgst")) + num(g("sgst")))
        else:
            tax = round(taxable * rate / 100.0, 2)
        key = norm_inv(raw_inv)
        rec = inv_map.setdefault(key, {
            "inv": raw_inv, "gstin": str(g("gstin") or "").strip().upper(),
            "party": str(g("party") or "").strip(), "date": g("date") or "",
            "taxable": 0.0, "tax": 0.0, "rates": set(), "lines": 0,
        })
        rec["taxable"] += taxable
        rec["tax"] += tax
        if rate:
            rec["rates"].add(round(rate, 2))
        rec["lines"] += 1
    print(f"  {label}: {source} -> {len(inv_map)} invoices "
          f"({'tax cols found' if have_tax else 'tax computed as taxable x rate'})")
    return inv_map

def gstin_key(rec, use_gstin):
    return (rec["gstin"] if use_gstin else "",)

# ---------------------------------------------------------------- G2/G3: reco
def reconcile(books, portal, tol):
    # G1: both sides have GSTIN on >50% of invoices -> include GSTIN in the key
    use_gstin = all(
        sum(1 for r in side.values() if len(r["gstin"]) >= 15) > 0.5 * max(1, len(side))
        for side in (books, portal))

    def keyset(side):
        return {(r["gstin"] if use_gstin else "") + "|" + k: k for k, r in side.items()}

    bk, pk = keyset(books), keyset(portal)
    matched_pairs = [(bk[k], pk[k]) for k in bk.keys() & pk.keys()]
    b_left = {v for v in bk.values()} - {b for b, _ in matched_pairs}
    p_left = {v for v in pk.values()} - {p for _, p in matched_pairs}

    # G1 pass 2: zeroless invoice numbers on the leftovers
    bz = {}
    for k in b_left:
        bz.setdefault((books[k]["gstin"] if use_gstin else "") + "|" + norm_inv_zeroless(k), k)
    for k in sorted(p_left):
        zk = (portal[k]["gstin"] if use_gstin else "") + "|" + norm_inv_zeroless(k)
        if zk in bz:
            matched_pairs.append((bz.pop(zk), k))
            p_left.discard(k)
    b_left = {k for k in bz.values()}

    matched, mismatch = [], []
    for bkey, pkey in matched_pairs:
        b, p = books[bkey], portal[pkey]
        d_taxable = round(b["taxable"] - p["taxable"], 2)
        d_tax = round(b["tax"] - p["tax"], 2)
        reasons = []
        if abs(d_taxable) > tol:
            reasons.append("TAXABLE_DIFF")
        if b["rates"] and p["rates"] and b["rates"] != p["rates"]:
            reasons.append("RATE_DIFF")          # G2: tolerance never forgives rate
        if abs(d_tax) > tol:
            reasons.append("TAX_DIFF")
        (mismatch if reasons else matched).append((b, p, d_taxable, d_tax, reasons))
    mismatch.sort(key=lambda t: -abs(t[3]))
    books_only = sorted((books[k] for k in b_left), key=lambda r: -r["tax"])
    portal_only = sorted((portal[k] for k in p_left), key=lambda r: -r["tax"])
    return matched, mismatch, books_only, portal_only, use_gstin

# ---------------------------------------------------------------- output
HDR_FILL = PatternFill("solid", start_color="1F4E79")
RED_FILL = PatternFill("solid", start_color="FFC7CE")
HDR_FONT = Font(bold=True, color="FFFFFF")

def sheet_with_header(wb, title, headers):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
    ws.freeze_panes = "A2"
    return ws

def rates_str(rates):
    return ", ".join(f"{r:g}%" for r in sorted(rates)) if rates else ""

def write_output(out_path, matched, mismatch, books_only, portal_only, tol, use_gstin):
    wb = Workbook()
    wb.remove(wb.active)

    tv = lambda rows: sum(r["taxable"] for r in rows)
    tx = lambda rows: sum(r["tax"] for r in rows)
    ws = sheet_with_header(wb, "SUMMARY", ["BUCKET", "INVOICES", "TAXABLE (books side)", "GST AMOUNT"])
    ws.append(["MATCHED", len(matched), round(tv([b for b, *_ in matched]), 2),
               round(tx([b for b, *_ in matched]), 2)])
    ws.append(["MISMATCH", len(mismatch), round(tv([b for b, *_ in mismatch]), 2),
               round(sum(abs(d) for *_, d, _r in [(m[0], m[1], m[3], m[4]) for m in mismatch]), 2)])
    ws.append(["BOOKS_ONLY (add to R1 — interest clock runs)", len(books_only),
               round(tv(books_only), 2), round(tx(books_only), 2)])
    ws.append(["PORTAL_ONLY (verify books / amend portal)", len(portal_only),
               round(tv(portal_only), 2), round(tx(portal_only), 2)])
    net = round(tx(books_only) - tx(portal_only)
                + sum(m[3] for m in mismatch), 2)
    ws.append([])
    ws.append(["NET GST AT STAKE (books minus portal; +ve = portal short)", "", "", net])
    ws.append([f"Tolerance ₹{tol:g}/invoice · match key = "
               f"{'GSTIN + ' if use_gstin else ''}normalized invoice no (G1)", "", "", ""])

    work_cols = ["ACTION", "AMEND_IN_PERIOD", "REMARKS"]

    ws = sheet_with_header(wb, "MISMATCH",
        ["INVOICE NO", "GSTIN", "PARTY", "DATE",
         "BOOKS TAXABLE", "PORTAL TAXABLE", "TAXABLE DIFF",
         "BOOKS RATE(S)", "PORTAL RATE(S)",
         "BOOKS GST", "PORTAL GST", "GST DIFF", "REASONS"] + work_cols)
    for b, p, dtv, dtx, reasons in mismatch:
        ws.append([b["inv"], b["gstin"] or p["gstin"], b["party"] or p["party"], str(b["date"]),
                   round(b["taxable"], 2), round(p["taxable"], 2), dtv,
                   rates_str(b["rates"]), rates_str(p["rates"]),
                   round(b["tax"], 2), round(p["tax"], 2), dtx,
                   " + ".join(reasons), "", "", ""])
        if "RATE_DIFF" in reasons or abs(dtx) > 100 * max(1, tol):
            for c in ws[ws.max_row]:
                c.fill = RED_FILL

    for title, rows in (("BOOKS_ONLY", books_only), ("PORTAL_ONLY", portal_only)):
        ws = sheet_with_header(wb, title,
            ["INVOICE NO", "GSTIN", "PARTY", "DATE", "TAXABLE", "RATE(S)", "GST AMOUNT"] + work_cols)
        for r in rows:
            ws.append([r["inv"], r["gstin"], r["party"], str(r["date"]),
                       round(r["taxable"], 2), rates_str(r["rates"]), round(r["tax"], 2),
                       "", "", ""])

    ws = sheet_with_header(wb, "MATCHED",
        ["INVOICE NO", "GSTIN", "PARTY", "DATE", "TAXABLE", "RATE(S)", "GST AMOUNT",
         "TAXABLE DIFF", "GST DIFF"])
    for b, p, dtv, dtx, _ in sorted(matched, key=lambda t: t[0]["inv"]):
        ws.append([b["inv"], b["gstin"], b["party"], str(b["date"]),
                   round(b["taxable"], 2), rates_str(b["rates"]), round(b["tax"], 2), dtv, dtx])

    for ws in wb.worksheets:
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(42, max(10, width + 2))
    wb.save(out_path)

# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="GSTR-1 books vs portal reconciliation (G1-G4)")
    ap.add_argument("books")
    ap.add_argument("portal")
    ap.add_argument("output")
    ap.add_argument("--tol", type=float, default=1.0, help="rounding tolerance per invoice (₹, default 1)")
    ap.add_argument("--books-sheet")
    ap.add_argument("--portal-sheet")
    a = ap.parse_args()

    print("Loading…")
    b_src, b_rows = read_rows(a.books, a.books_sheet)
    p_src, p_rows = read_rows(a.portal, a.portal_sheet)
    books = load_side("BOOKS ", b_src, b_rows)
    portal = load_side("PORTAL", p_src, p_rows)

    matched, mismatch, books_only, portal_only, use_gstin = reconcile(books, portal, a.tol)
    write_output(a.output, matched, mismatch, books_only, portal_only, a.tol, use_gstin)

    print(f"\nRECO RESULT (tol ₹{a.tol:g}, key={'GSTIN+' if use_gstin else ''}inv no):")
    print(f"  MATCHED     : {len(matched)}")
    print(f"  MISMATCH    : {len(mismatch)}")
    print(f"  BOOKS_ONLY  : {len(books_only)}  (add to R1 — interest clock)")
    print(f"  PORTAL_ONLY : {len(portal_only)}  (verify books / amend)")
    print(f"  -> {a.output}")

if __name__ == "__main__":
    main()
