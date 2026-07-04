#!/usr/bin/env python3
"""FY25-26 monthly worklist (M17/M18 adapted for M13_FINAL layout).
Computes recovery/ESI/wage-code universes from scratch (no ACTION_NEEDED flags).
Usage: python fy2526_worklist.py <Month_M13_FINAL.xlsx> <out.xlsx> <MonthName>"""
import sys, re
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill

HDR_FILL = PatternFill("solid", start_color="1F4E79"); HDR_FONT = Font(bold=True, color="FFFFFF")
HI_FILL = PatternFill("solid", start_color="FFC7CE")

def num(v):
    if v is None: return 0.0
    if isinstance(v,(int,float)): return float(v)
    s = re.sub(r"[^0-9.\-]","",str(v))
    try: return float(s) if s not in ("","-",".") else 0.0
    except ValueError: return 0.0

def style_header(ws,n):
    for c in range(1,n+1):
        ws.cell(row=ws.max_row,column=c).fill=HDR_FILL; ws.cell(row=ws.max_row,column=c).font=HDR_FONT
    ws.freeze_panes="A2"

def autowidth(ws,w):
    for i,x in enumerate(w,1): ws.column_dimensions[ws.cell(row=1,column=i).column_letter].width=x

src, dst, MON = sys.argv[1], sys.argv[2], sys.argv[3]
wb = load_workbook(src, read_only=True, data_only=True)
ws = wb["Salary (Corrected)"]
it = ws.iter_rows(values_only=True)
hdr = None
for row in it:
    if row and any(str(v).strip()=="EMPCODE" for v in row if v): hdr=list(row); break
ix = {str(h).strip():i for i,h in enumerate(hdr) if h is not None}
rows = [r for r in it if r is not None and any(v is not None for v in r)]
def g(r,c): return r[ix[c]] if c in ix and ix[c]<len(r) else None
def gn(r,c): return num(g(r,c))
RGN = 'REVISED_GROSS_NEW (final)' if 'REVISED_GROSS_NEW (final)' in ix else 'REVISED_GROSS_NEW'
print(f"{MON}: {len(rows)} rows, {len(ix)} cols; DA={'DA' in ix} CTC={'CTC' in ix} OT_AMT={'OT AMOUNT' in ix}")

# ---- golden-rule check ----
tot = lambda c: sum(gn(r,c) for r in rows)
net, rnet = tot('NETPAYABLE'), tot('REVISED_NET_PAYABLE')
pf, ecr = tot('REVISED_PF'), tot('ECR_PF')
print(f"  GOLDEN: NET {net:,.0f} vs REVISED_NET {rnet:,.0f} (drift {rnet-net:,.0f}) | REVISED_PF {pf:,.0f} vs ECR {ecr:,.0f} (gap {pf-ecr:,.0f})")

VAR_COLS = ['OT AMOUNT','EXTRA OT','BASIC DA ARREARS','OTHER ARREARS','ATTENDANCE ALW',
            'PAID LEAVE','NATIONAL /FESTIVAL HOLIDAY','FESTIVAL HOLIDAY','WEEKLY OFF',
            'TRAVELLING ALLOWANCE','PAID HOLIDAY','NATIONAL HOLIDAY']
out = Workbook()

# ---- RECOVERY_REVIEW (M18 computed on every paid row) ----
rec_cols = ['PRIORITY','EMPCODE','FULLNAME','CLIENTGROUPNAME','SITENAME','SITESTATE',
            'DESIGNATIONNAME','RATE_BASIS','FIXEDGROSS','NORMALDAYS',
            'EXPECTED_BASE (rate/div×days)','ALLOWANCES (OT/arr/att/hol)','EXPECTED_TOTAL',
            'GROSS AMT','REVISED_GROSS','EXCESS_TO_REVIEW','VERIFIED (Y/N)',
            'DECISION (RECOVER/WAIVE/JUSTIFIED)','RECOVERY_MONTH','REMARKS']
wr = out.active; wr.title="RECOVERY_REVIEW"; wr.append(rec_cols); style_header(wr,len(rec_cols))
kept=[]; cleared=0; scanned=0
for r in rows:
    fg, gr, days = gn(r,'FIXEDGROSS'), gn(r,'GROSS AMT'), gn(r,'NORMALDAYS')
    if fg<=0 or gr<=0 or days<=0: continue
    scanned+=1
    div = gn(r,'SITEDIVISIONDAYS') or 30
    var = sum(max(0.0,gn(r,c)) for c in VAR_COLS if c in ix)
    cand = {'PER-DAY': fg*days+var, f'MONTHLY/{int(div)}': (fg/div*days if div>0 else fg)+var}
    if div==1 or fg<3000: basis='PER-DAY'
    elif fg<6000: basis=min(cand,key=lambda k:abs(gr-cand[k]))
    else: basis=f'MONTHLY/{int(div)}'
    expected=cand[basis]; base=expected-var
    # M18d (M13-FINAL layout): reconciliation already restructured gross holding
    # NET — original GROSS AMT may carry components clawed back via deductions.
    # Measure excess against the MOST CONSERVATIVE positive gross figure.
    grosses=[c_ for c_ in (gr,gn(r,'REVISED_GROSS'),gn(r,RGN)) if c_>0]
    gmin=min(grosses)
    exc=max(0.0,min(gmin-expected,gmin))
    # M18e cash cap: recoverable excess cannot exceed the row's actual take-home
    # (M13 clawback rows carry inflated gross offset by deductions, NET ~ small)
    exc=min(exc,max(0.0,gn(r,'NETPAYABLE')))
    tol=max(500.0,expected*0.02)
    close=any(abs(c_-expected)<=tol for c_ in grosses)
    if exc<1 or close: cleared+=1; continue
    kept.append((exc,basis,days,fg,gr,base,var,expected,gn(r,'REVISED_GROSS'),r))
kept.sort(key=lambda x:-x[0])
n_hi=0; rec_total=hi_total=0.0
for exc,basis,days,fg,gr,base,var,expected,rg,r in kept:
    hi = exc>10000; n_hi+=hi; rec_total+=exc; hi_total+=exc if hi else 0
    wr.append(['HIGH' if hi else 'LOW',g(r,'EMPCODE'),g(r,'FULLNAME'),g(r,'CLIENTGROUPNAME'),
               g(r,'SITENAME'),g(r,'SITESTATE'),g(r,'DESIGNATIONNAME'),basis,round(fg,0),days,
               round(base,0),round(var,0),round(expected,0),round(gr,0),round(rg,0),round(exc,0),'','','',''])
    if hi:
        for c in range(1,len(rec_cols)+1): wr.cell(row=wr.max_row,column=c).fill=HI_FILL
autowidth(wr,[8,11,24,26,24,14,18,11,10,8,13,12,12,10,11,12,9,16,12,24])

# ---- ESI_ENROLLMENT (computed: ESIC=0 & full-month rate <= 21,000) ----
esi_cols=['EMPCODE','FULLNAME','CLIENTGROUPNAME','SITENAME','SITESTATE','ESIC NO','UAN NO',
          'GROSS AMT','REAL_FULL_MONTH_GROSS','ESI_EE_0.75%','ESI_ER_3.25%','IC_NO_ALLOTTED',
          'ENROLLED_FROM (month)','REMARKS']
we=out.create_sheet("ESI_ENROLLMENT"); we.append(esi_cols); style_header(we,len(esi_cols))
esi=[]
for r in rows:
    fg, gr, days = gn(r,'FIXEDGROSS'), gn(r,'GROSS AMT'), gn(r,'NORMALDAYS')
    if fg<=0 or gr<=0 or days<=0: continue
    if gn(r,'ESIC')>0 or gn(r,'REVISED_ESIC')>0: continue     # already contributing
    div = gn(r,'SITEDIVISIONDAYS') or 30
    full = fg*26 if (div==1 or fg<3000) else fg               # M18c basis
    if 0<full<=21000: esi.append((full,gr,r))
esi.sort(key=lambda t:(str(g(t[2],'CLIENTGROUPNAME') or ''),-t[1]))
esi_gross=0.0
for full,gr,r in esi:
    esi_gross+=gr
    we.append([g(r,'EMPCODE'),g(r,'FULLNAME'),g(r,'CLIENTGROUPNAME'),g(r,'SITENAME'),
               g(r,'SITESTATE'),g(r,'ESIC NO'),g(r,'UAN NO'),round(gr,0),round(full,0),
               round(gr*0.0075,0),round(gr*0.0325,0),'','',''])
autowidth(we,[11,24,26,24,14,14,14,10,12,10,10,14,12,24])

# ---- WAGE_CODE_50PCT ----
wc_cols=['SEVERITY','NOTE','EMPCODE','FULLNAME','CLIENTGROUPNAME','SITENAME','SITESTATE',
         'DESIGNATIONNAME','NORMALDAYS','BASIC','DA','BASIC+DA','OT+ARREARS','GROSS AMT','CTC',
         'BASIC+DA % OF CTC','SHORTFALL_TO_50PCT','RESTRUCTURE (Y/N)','TARGET_BASIC_DA (=50% CTC)','REMARKS']
ww=out.create_sheet("WAGE_CODE_50PCT"); ww.append(wc_cols); style_header(ww,len(wc_cols))
wage_fail=[]; neg_gross=0
for r in rows:
    b,d,ctc,gr = gn(r,'BASIC'),gn(r,'DA'),gn(r,'CTC'),gn(r,'GROSS AMT')
    if ctc<=0 or (b+d)<=0: continue
    if gr<0: neg_gross+=1; continue
    ratio=(b+d)/ctc
    if ratio<0.5:
        ot_arr=gn(r,'OT AMOUNT')+gn(r,'EXTRA OT')+gn(r,'BASIC DA ARREARS')+gn(r,'OTHER ARREARS')
        wage_fail.append((ratio,b,d,ctc,gr,ot_arr,r))
wage_fail.sort(key=lambda x:x[0])
n_sev=n_struct=0
for ratio,b,d,ctc,gr,ot_arr,r in wage_fail:
    sev='<30%' if ratio<0.30 else ('30-40%' if ratio<0.40 else ('40-45%' if ratio<0.45 else '45-50%'))
    note=''
    if ot_arr>gr*0.5: note='OT/ARREARS-HEAVY — remuneration inflated this month'
    elif gn(r,'NORMALDAYS')<=2: note='LOW-DAY ROW — verify'
    else: n_struct+=1
    target=round(ctc*0.5,0)
    ww.append([sev,note,g(r,'EMPCODE'),g(r,'FULLNAME'),g(r,'CLIENTGROUPNAME'),g(r,'SITENAME'),
               g(r,'SITESTATE'),g(r,'DESIGNATIONNAME'),gn(r,'NORMALDAYS'),round(b,0),round(d,0),
               round(b+d,0),round(ot_arr,0),round(gr,0),round(ctc,0),round(ratio*100,1),
               round(target-(b+d),0),'',target,''])
    if ratio<0.30 and not note:
        n_sev+=1
        for c in range(1,len(wc_cols)+1): ww.cell(row=ww.max_row,column=c).fill=HI_FILL
autowidth(ww,[8,34,11,24,26,24,14,18,8,9,8,10,10,10,10,11,12,9,14,24])

# ---- SUMMARY ----
s=out.create_sheet("SUMMARY",0)
s.append([f"{MON} 2025 WORKLIST — RECOVERY / ESI / WAGE-CODE (M17+M18, computed from M13 FINAL)"])
s.cell(row=1,column=1).font=Font(bold=True,size=13)
s.append([f"Source: {src.split('/')[-1]} · golden rules: NET drift {rnet-net:,.0f}, PF-ECR gap {pf-ecr:,.0f}"])
s.append([])
s.append(["WORKSTREAM","EMPLOYEES","AMOUNT (Rs.)","WHERE"]); 
for c in range(1,5): s.cell(row=4,column=c).fill=HDR_FILL; s.cell(row=4,column=c).font=HDR_FONT
s.append(["Recovery review — HIGH (> Rs.10k each)",n_hi,round(hi_total,0),"Tab RECOVERY_REVIEW (red, on top)"])
s.cell(row=5,column=1).fill=HI_FILL
s.append(["Recovery review — remaining small rows",len(kept)-n_hi,round(rec_total-hi_total,0),"Tab RECOVERY_REVIEW"])
s.append(["Rows cleared by M18 (expected ≈ gross within max(500,2%))",cleared,"","not shown — no overpayment"])
s.append(["ESI enrollment needed (exempt but rate ≤ 21k)",len(esi),round(esi_gross*0.04,0),"Tab ESI_ENROLLMENT (4%/month exposure)"])
wc_short=sum(max(0.0,ctc*0.5-(b+d)) for _,b,d,ctc,_g,_o,_r in wage_fail)
s.append([f"Wage Code 50% FAIL ({n_struct} structural, {n_sev} severe <30%; {neg_gross} negative-gross excluded)",
          len(wage_fail),round(wc_short,0),"Tab WAGE_CODE_50PCT (shortfall to 50%)"])
autowidth(s,[64,11,14,60])
out.save(dst)
print(f"  RECOVERY: {len(kept):,} rows Rs.{rec_total:,.0f} (HIGH {n_hi} Rs.{hi_total:,.0f}; cleared {cleared:,} of {scanned:,} scanned)")
print(f"  ESI     : {len(esi):,} rows, exposure Rs.{esi_gross*0.04:,.0f}/month")
print(f"  WAGECODE: {len(wage_fail):,} fails ({n_struct} structural, {n_sev} severe) shortfall Rs.{wc_short:,.0f}")
print(f"  -> {dst}")
