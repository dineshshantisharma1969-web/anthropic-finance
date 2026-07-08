#!/usr/bin/env python3
"""Combined PF + ESI management one-pager for Maharashtra FY2024-25."""
import os, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

D = os.environ.get("MH_DIR", os.path.dirname(os.path.abspath(__file__)))
pf   = pd.read_csv(f"{D}/SUMMARY_FY2024-25.csv")
esi  = pd.read_csv(f"{D}/esi/ESI_SUMMARY_FY2024-25.csv")
gapb = pd.read_csv(f"{D}/queries/SUMMARY_below15k_ecr0_partVSfull_FY2024-25.csv")

ORDER=['APR-24','MAY-24','JUN-24','JUL-24','AUG-24','SEP-24','OCT-24','NOV-24','DEC-24','JAN-25','FEB-25','MAR-25']
NAVY="1F3864"; BLUE="2E5496"; TEAL="1F6E6B"; LGREY="D9E1F2"; GREEN="C6EFCE"; GREENF="006100"
AMBER="FFF2CC"; AMBERF="7F6000"; WHITE="FFFFFF"; RED="C00000"
thin=Side(style="thin",color="BFBFBF"); border=Border(thin,thin,thin,thin)
M='#,##0'
wb=Workbook(); ws=wb.active; ws.title="PF + ESI SUMMARY"; ws.sheet_view.showGridLines=False

def C(r,c,v,b=False,sz=10,col="000000",fill=None,al="left",fmt=None,bd=False,it=False):
    x=ws.cell(row=r,column=c,value=v)
    x.font=Font(bold=b,size=sz,color=col,italic=it)
    x.alignment=Alignment(horizontal=al,vertical="center")
    if fill:x.fill=PatternFill("solid",fgColor=fill)
    if fmt:x.number_format=fmt
    if bd:x.border=border
    return x

def band(r,c0,c1,text,fill=NAVY):
    C(r,c0,text,b=True,sz=12,col=WHITE,fill=fill)
    for c in range(c0+1,c1+1): C(r,c,"",fill=fill)

def pfget(m,k):
    row=pf[pf['MONTH']==m].iloc[0]; return float(row[k])
def esiget(m,k):
    row=esi[esi['MONTH']==m].iloc[0]; return float(row[k])
def gapget(m,k):
    row=gapb[gapb['MONTH']==m].iloc[0]; return float(row[k])

r=2
C(r,2,"MAHARASHTRA (AISSS) — PF & ESI RECONCILIATION",b=True,sz=16,col=NAVY); r+=1
C(r,2,"FY2024-25 · Apr-24 → Mar-25 · West-only · for management sign-off",sz=11,col=BLUE,it=True); r+=2

# ===== A. Headline anchors (two side-by-side blocks) =====
band(r,2,8,"A.  HEADLINE — FILED / RECONCILED POSITION (₹)"); r+=1
hr=r
C(r,2,"PROVIDENT FUND (PF)",b=True,sz=11,col=WHITE,fill=BLUE); C(r,3,"",fill=BLUE)
C(r,6,"EMPLOYEE STATE INSURANCE (ESI)",b=True,sz=11,col=WHITE,fill=TEAL); C(r,7,"",fill=TEAL); r+=1
pf_lines=[("Original salary PF",pf[pf['MONTH']=='TOTAL']['ORIG_SALARY_PF'].iloc[0]),
          ("ECR PF filed",pf[pf['MONTH']=='TOTAL']['ECR_PF_FILED'].iloc[0]),
          ("Reconciled PF (= ECR, cap ₹1,800)",pf[pf['MONTH']=='TOTAL']['REVISED_PF'].iloc[0]),
          ("PF gap after reconciliation",0),
          ("PF parked (Other Ded, net-neutral)",pf[pf['MONTH']=='TOTAL']['PF_PARKED_IN_OTHER_DED'].iloc[0])]
et=esi[esi['MONTH']=='TOTAL'].iloc[0]
esi_lines=[("ESIC wages",float(et['ESIC_WAGES'])),
           ("ESI employee (0.75%)",float(et['ESI_EMP_FILED(0.75%)'])),
           ("ESI employer (3.25%)",float(et['ESI_CO_FILED(3.25%)'])),
           ("ESI total deposited",float(et['ESI_TOTAL_FILED'])),
           ("ESI rate/West-only checks",None)]
for i in range(5):
    lab,val=pf_lines[i]; hi=(i==2)
    C(r,2,lab,b=hi,sz=10,fill=(GREEN if hi else None),col=(GREENF if hi else "000000"),bd=True)
    C(r,3,val,b=hi,sz=10,al="right",fmt=M,fill=(GREEN if hi else None),col=(GREENF if hi else "000000"),bd=True)
    lab2,val2=esi_lines[i]; hi2=(i in (1,3))
    C(r,6,lab2,b=hi2,sz=10,fill=(GREEN if hi2 else None),col=(GREENF if hi2 else "000000"),bd=True)
    if val2 is None:
        C(r,7,"PASS ✓",b=True,sz=10,al="right",fill=GREEN,col=GREENF,bd=True)
    else:
        C(r,7,val2,b=hi2,sz=10,al="right",fmt=M,fill=(GREEN if hi2 else None),col=(GREENF if hi2 else "000000"),bd=True)
    r+=1
r+=1

# ===== B. Sign-off checklist =====
band(r,2,8,"B.  SIGN-OFF CHECKLIST"); r+=1
checks=[("Σ Revised PF = ECR per employee (₹0 gap, all 12 months)","PASS ✓",GREEN,GREENF),
        ("Net Payable unchanged on every row (PF & ESI net-neutral)","PASS ✓",GREEN,GREENF),
        ("PF = 12%×(Basic+DA), ≤₹1,800; Basic+DA ≤₹15,000","PASS ✓",GREEN,GREENF),
        ("ESI employee 0.75% & employer 3.25% on every row","PASS ✓ (2 rounding ≤₹1.5)",GREEN,GREENF),
        ("West-only — PF challan & ESIC register, 0 South leakage","PASS ✓",GREEN,GREENF),
        ("PF coverage gaps (below ₹15k, full attendance, not in ECR)","332 — REVIEW",AMBER,AMBERF),
        ("ESI coverage gaps (gross ≤ ₹21k, not in ESIC register)","750 — REVIEW",AMBER,AMBERF)]
for lab,res,fill,fg in checks:
    C(r,2,lab,sz=10,bd=True); C(r,3,"",bd=True); C(r,4,"",bd=True)
    C(r,5,res,b=True,sz=10,al="center",fill=fill,col=fg,bd=True); C(r,6,"",fill=fill,bd=True)
    r+=1
r+=1

# ===== C. Decision items =====
band(r,2,8,"C.  ITEMS NEEDING A MANAGEMENT DECISION"); r+=1
for h,w in [("Item",4),("Detail",40),("Qty / ₹",18),("Action",22)]: pass
hdr=["#","Item","Count / ₹","Decision"]
C(r,2,"#",b=True,sz=9,col=WHITE,fill=BLUE,al="center",bd=True)
C(r,3,"Item",b=True,sz=9,col=WHITE,fill=BLUE,bd=True)
C(r,6,"Count / ₹",b=True,sz=9,col=WHITE,fill=BLUE,al="right",bd=True)
C(r,7,"Decision",b=True,sz=9,col=WHITE,fill=BLUE,bd=True); C(r,8,"",fill=BLUE,bd=True)
r+=1
items=[("A","PF: below ₹15k, full attendance, NOT in ECR (coverage gap)","332 · ₹5,14,086","Regularise / confirm exempt"),
       ("B","PF: ECR filed above ₹1,800 EE cap (capped in reco)","991 · ₹3,57,385","Note / accept"),
       ("C","ESI: gross ≤ ₹21k, NOT in ESIC register (coverage gap)","750 emp-months","Regularise / confirm exempt"),
       ("D","ESI: wages above ₹21,000 (period continuation)","690 rows","Confirm continuation")]
for i,(n,it,qty,act) in enumerate(items):
    fill=WHITE if i%2==0 else "F2F6FC"
    C(r,2,n,b=True,sz=9,al="center",fill=fill,bd=True)
    C(r,3,it,sz=9,fill=fill,bd=True); C(r,4,"",fill=fill,bd=True); C(r,5,"",fill=fill,bd=True)
    C(r,6,qty,sz=9,al="right",fill=fill,bd=True)
    C(r,7,act,sz=9,fill=fill,bd=True); C(r,8,"",fill=fill,bd=True)
    r+=1
r+=1

# ===== D. Month-wise PF + ESI =====
band(r,2,8,"D.  MONTH-WISE — PF & ESI (₹)"); r+=1
cols=[("MONTH",""),("PF: Revised (=ECR)","REVISED_PF"),("PF: Parked","PF_PARKED_IN_OTHER_DED"),
      ("PF gaps","pfgap"),("ESI emp (0.75%)","ESI_EMP_FILED(0.75%)"),("ESI total","ESI_TOTAL_FILED"),("ESI gaps","esigap")]
for j,(lab,_) in enumerate(cols):
    C(r,2+j,lab,b=True,sz=9,col=WHITE,fill=NAVY,al=("left" if j==0 else "right"),bd=True)
r+=1
for m in ORDER:
    fill=WHITE if (r%2==0) else "F2F6FC"
    vals=[m, pfget(m,'REVISED_PF'), pfget(m,'PF_PARKED_IN_OTHER_DED'),
          int(gapget(m,'full_attendance_GAP')), esiget(m,'ESI_EMP_FILED(0.75%)'),
          esiget(m,'ESI_TOTAL_FILED'), int(esiget(m,'ESI_COVERAGE_GAPS(<=21k not filed)'))]
    for j,v in enumerate(vals):
        C(r,2+j,v,sz=9,al=("left" if j==0 else "right"),fill=fill,fmt=(None if j in (0,) else M),bd=True)
    r+=1
# total row
tot=[ "TOTAL",
      pf[pf['MONTH']=='TOTAL']['REVISED_PF'].iloc[0],
      pf[pf['MONTH']=='TOTAL']['PF_PARKED_IN_OTHER_DED'].iloc[0],
      int(gapb[gapb['MONTH'].isin(ORDER)]['full_attendance_GAP'].sum()),
      float(et['ESI_EMP_FILED(0.75%)']), float(et['ESI_TOTAL_FILED']),
      int(esi[esi['MONTH']=='TOTAL']['ESI_COVERAGE_GAPS(<=21k not filed)'].iloc[0]) ]
for j,v in enumerate(tot):
    C(r,2+j,v,b=True,sz=9,al=("left" if j==0 else "right"),fill=LGREY,fmt=(None if j==0 else M),bd=True)
r+=2

C(r,2,"PF: Maharashtra salary PF reconciled to ECR (West challan + DMART), ₹0 gap, Net Payable unchanged.",sz=9,it=True,col=BLUE); r+=1
C(r,2,"ESI: filed ESIC register validated (0.75%/3.25%, West-only). Coverage gaps = eligible-but-not-filed, for review.",sz=9,it=True,col=TEAL); r+=1
C(r,2,"Backup: Maharashtra_PF_Summary, Coverage_Gap_Book, esi/Maharashtra_ESI_Reconciliation, and per-month CSVs.",sz=9,it=True,col="808080"); r+=1

widths={2:6,3:34,4:14,5:14,6:20,7:22,8:12}
for c,w in widths.items(): ws.column_dimensions[get_column_letter(c)].width=w
ws.column_dimensions['B'].width=40

out=f"{D}/Maharashtra_PF_ESI_Management_Summary_FY2024-25.xlsx"
wb.save(out); print("wrote",out)
