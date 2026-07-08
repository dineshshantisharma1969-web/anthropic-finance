#!/usr/bin/env python3
"""Build a single year-wide PF coverage-gap workbook for Maharashtra FY2024-25:
employees with Basic(+DA) < 15,000 and ECR PF = 0, split part-days vs full-attendance."""
import os, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

D = os.environ.get("MH_DIR", os.path.dirname(os.path.abspath(__file__)))
CM = f"{D}/corrected_monthly"
ORDER=['APR-24','MAY-24','JUN-24','JUL-24','AUG-24','SEP-24','OCT-24','NOV-24','DEC-24','JAN-25','FEB-25','MAR-25']
FULL_MONTH={'APR-24':30,'MAY-24':31,'JUN-24':30,'JUL-24':31,'AUG-24':31,'SEP-24':30,
            'OCT-24':31,'NOV-24':30,'DEC-24':31,'JAN-25':31,'FEB-25':28,'MAR-25':31}

frames=[]
for m in ORDER:
    d=pd.read_csv(f"{CM}/CORRECTED_{m}.csv")
    ecr=pd.to_numeric(d['ECR_PF (EE, filed)'],errors='coerce').fillna(0)
    basic=pd.to_numeric(d['ORIG_BASIC(+DA)'],errors='coerce').fillna(0)
    mdays=pd.to_numeric(d['M/DAYS (divisor)'],errors='coerce').fillna(0)
    wdays=pd.to_numeric(d['NORMALDAYS (W/DAYS)'],errors='coerce').fillna(0)
    ncp=pd.to_numeric(d['NCP'],errors='coerce').fillna(0)
    sel=(ecr==0)&(basic<15000)
    g=d[sel].copy()
    g.insert(0,'MONTH',m)
    part=(wdays[sel]<mdays[sel])|(ncp[sel]>0)
    g['ATTENDANCE']=part.map({True:'PART_DAYS',False:'FULL_ATTENDANCE_GAP'}).values
    # fixed monthly basic = daily rate x full standard month (divisor) — the contracted monthly wage
    dr=(basic[sel]/wdays[sel].replace(0,pd.NA))
    g['FIXED_MONTHLY_BASIC']=(dr*mdays[sel]).round(0).values
    frames.append(g)
allg=pd.concat(frames,ignore_index=True)

cols=['MONTH','EMP CODE','NAME','DESIGNATION','BRANCH','SITE NAME','M/DAYS (divisor)',
      'NORMALDAYS (W/DAYS)','NCP','ORIG_BASIC(+DA)','FIXED_MONTHLY_BASIC','ORIG_PF (salary EE)',
      'REVISED_PF','BD_MONTHLY_PROJECTION','PF_DIFF_PARKED_IN_OTHER_DED','ATTENDANCE']
allg=allg[[c for c in cols if c in allg.columns]]
full=allg[allg['ATTENDANCE']=='FULL_ATTENDANCE_GAP'].copy()

# ---- styling ----
NAVY="1F3864"; BLUE="2E5496"; LGREY="D9E1F2"; GREEN="C6EFCE"; GREENF="006100"; RED="C00000"
REDF="9C0006"; REDFILL="FCE4E4"; AMBER="FFF2CC"; WHITE="FFFFFF"
thin=Side(style="thin",color="BFBFBF"); border=Border(thin,thin,thin,thin)
money='#,##0'
wb=Workbook()

def style_header(ws,row,headers,fill=BLUE):
    for j,h in enumerate(headers):
        c=ws.cell(row=row,column=1+j,value=h)
        c.font=Font(bold=True,size=9,color=WHITE); c.fill=PatternFill("solid",fgColor=fill)
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=border

def write_df(ws,df,start,money_cols=(),int_cols=(),highlight_gap=True):
    style_header(ws,start,list(df.columns))
    r=start+1
    for _,row in df.iterrows():
        is_gap = str(row.get('ATTENDANCE',''))=='FULL_ATTENDANCE_GAP'
        base_fill = (REDFILL if (highlight_gap and is_gap) else (WHITE if r%2==0 else "F2F6FC"))
        for j,col in enumerate(df.columns):
            v=row[col]
            fmt=money if (col in money_cols or col in int_cols) else None
            c=ws.cell(row=r,column=1+j,value=(int(v) if col in int_cols and pd.notna(v) else v))
            c.font=Font(size=9); c.fill=PatternFill("solid",fgColor=base_fill)
            c.alignment=Alignment(horizontal=("right" if fmt else "left"),vertical="center")
            if fmt:c.number_format=fmt
            c.border=border
        r+=1
    return r

MONEY={'ORIG_BASIC(+DA)','FIXED_MONTHLY_BASIC','ORIG_PF (salary EE)','REVISED_PF','BD_MONTHLY_PROJECTION','PF_DIFF_PARKED_IN_OTHER_DED'}
INTS={'M/DAYS (divisor)','NORMALDAYS (W/DAYS)','NCP'}

# ===== Sheet 1: SUMMARY =====
ws=wb.active; ws.title="SUMMARY"; ws.sheet_view.showGridLines=False
def cell(r,c,v,**k):
    x=ws.cell(row=r,column=c,value=v)
    x.font=Font(bold=k.get('bold',False),size=k.get('size',10),color=k.get('color','000000'),italic=k.get('italic',False))
    x.alignment=Alignment(horizontal=k.get('align','left'),vertical="center")
    if k.get('fill'):x.fill=PatternFill("solid",fgColor=k['fill'])
    if k.get('fmt'):x.number_format=k['fmt']
    if k.get('border'):x.border=border
    return x
cell(1,1,"MAHARASHTRA (AISSS) — PF COVERAGE-GAP BOOK · FY2024-25",bold=True,size=15,color=NAVY)
cell(2,1,"Employees with Basic(+DA) < ₹15,000 AND ECR PF = 0  ·  West-only reconciliation",size=11,color=BLUE,italic=True)
# month-wise table
r=4
hdr=['MONTH','Below-15k & ECR=0','Part-days','FULL-ATTENDANCE GAP','PF Parked (part)','PF Parked (GAP)']
style_header(ws,r,hdr,fill=NAVY); r+=1
tot={'below':0,'part':0,'gap':0,'ppart':0.0,'pgap':0.0}
for m in ORDER:
    sub=allg[allg['MONTH']==m]; g=sub[sub['ATTENDANCE']=='FULL_ATTENDANCE_GAP']; p=sub[sub['ATTENDANCE']=='PART_DAYS']
    vals=[m,len(sub),len(p),len(g),round(p['PF_DIFF_PARKED_IN_OTHER_DED'].sum(),0),round(g['PF_DIFF_PARKED_IN_OTHER_DED'].sum(),0)]
    tot['below']+=len(sub);tot['part']+=len(p);tot['gap']+=len(g);tot['ppart']+=vals[4];tot['pgap']+=vals[5]
    for j,v in enumerate(vals):
        cell(r,1+j,v,size=9,align=("left" if j==0 else "right"),fill=(WHITE if r%2==0 else "F2F6FC"),
             fmt=(money if j in(4,5) else None),border=True,
             color=(REDF if j==3 and len(g)>0 else '000000'),bold=(j==3 and len(g)>0))
    r+=1
totvals=['TOTAL',tot['below'],tot['part'],tot['gap'],tot['ppart'],tot['pgap']]
for j,v in enumerate(totvals):
    cell(r,1+j,v,bold=True,size=9,align=("left" if j==0 else "right"),fill=LGREY,fmt=(money if j in(4,5) else None),border=True)
r+=2
cell(r,1,f"GAP = full-attendance employees (worked the full standard month) whose monthly basic is still below ₹15,000 yet not filed in ECR — PF coverage-review candidates.",size=9,italic=True,color=BLUE); r+=1
cell(r,1,f"Total gap rows: {tot['gap']}  ·  PF on gap rows: ₹{tot['pgap']:,.0f}  ·  see 'GAP_FULL_ATTENDANCE' tab for the line list.",size=9,italic=True,color=RED); r+=1
cell(r,1,"Part-days rows: low basic mainly due to partial attendance (reconciliation projects them above the ceiling). See 'ALL_BELOW15K_ECR0'.",size=9,italic=True,color="808080"); r+=1
# branch pivot of gaps
r+=1
cell(r,1,"GAP ROWS BY BRANCH",bold=True,size=11,color=WHITE,fill=NAVY); cell(r,2,"",fill=NAVY); r+=1
bp=full.groupby('BRANCH').agg(rows=('EMP CODE','size'),pf=('PF_DIFF_PARKED_IN_OTHER_DED','sum')).reset_index()
style_header(ws,r,['BRANCH','Gap rows','PF on gap rows']); r+=1
for _,row in bp.iterrows():
    cell(r,1,row['BRANCH'],size=9,border=True,fill=(WHITE if r%2 else "F2F6FC"))
    cell(r,2,int(row['rows']),size=9,align="right",border=True,fill=(WHITE if r%2 else "F2F6FC"))
    cell(r,3,round(row['pf'],0),size=9,align="right",fmt=money,border=True,fill=(WHITE if r%2 else "F2F6FC"))
    r+=1
for col,w in {1:22,2:20,3:22,4:22,5:18,6:18}.items(): ws.column_dimensions[get_column_letter(col)].width=w

# ===== Sheet 2: GAP_FULL_ATTENDANCE =====
ws=wb.create_sheet("GAP_FULL_ATTENDANCE"); ws.sheet_view.showGridLines=False
c=ws.cell(row=1,column=1,value="FULL-ATTENDANCE below-₹15,000 not-in-ECR — PF coverage-gap candidates (FY2024-25)")
c.font=Font(bold=True,size=12,color=RED)
write_df(ws,full.sort_values(['MONTH','BRANCH','EMP CODE']),3,money_cols=MONEY,int_cols=INTS)
ws.freeze_panes="A4"
for j,col in enumerate(full.columns):
    ws.column_dimensions[get_column_letter(1+j)].width=max(12,min(40,int(full[col].astype(str).str.len().max())+2))

# ===== Sheet 3: ALL_BELOW15K_ECR0 =====
ws=wb.create_sheet("ALL_BELOW15K_ECR0"); ws.sheet_view.showGridLines=False
c=ws.cell(row=1,column=1,value="ALL below-₹15,000 not-in-ECR (part-days + full-attendance), FY2024-25")
c.font=Font(bold=True,size=12,color=NAVY)
write_df(ws,allg.sort_values(['MONTH','ATTENDANCE','BRANCH','EMP CODE']),3,money_cols=MONEY,int_cols=INTS)
ws.freeze_panes="A4"
for j,col in enumerate(allg.columns):
    ws.column_dimensions[get_column_letter(1+j)].width=max(12,min(40,int(allg[col].astype(str).str.len().max())+2))

out=f"{D}/Maharashtra_PF_Coverage_Gap_Book_FY2024-25.xlsx"
wb.save(out)
# also a flat CSV of the gap set
full.sort_values(['MONTH','BRANCH','EMP CODE']).to_csv(f"{D}/queries/YEAR_full_attendance_GAP_ecr0_FY2024-25.csv",index=False)
print("wrote",out)
print(f"gap rows={len(full)}  all-below15k rows={len(allg)}  gap PF=Rs.{full['PF_DIFF_PARKED_IN_OTHER_DED'].sum():,.0f}")
print("branches:",bp.to_dict('records'))
