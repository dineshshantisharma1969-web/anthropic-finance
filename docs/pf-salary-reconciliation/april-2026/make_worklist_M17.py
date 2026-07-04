#!/usr/bin/env python3
"""
Rule M17 — Monthly Recovery & ESI Worklist generator.

Takes a reconciled monthly file (AFTER Rule M16 is applied) and produces
<Mon>_WORKLIST_Recovery_ESI.xlsx with three tabs:

  1. SUMMARY          — headline counts/amounts + by-client breakdown
  2. RECOVERY_REVIEW  — overpaid-vs-rate employees, HIGH priority (> Rs.10k)
                        highlighted & on top, with working columns:
                        VERIFIED / DECISION / RECOVERY_MONTH / REMARKS
  3. ESI_ENROLLMENT   — ESI-eligible-but-exempt employees grouped by client,
                        with EE/ER contribution columns and working columns:
                        IC_NO_ALLOTTED / ENROLLED_FROM / REMARKS

Usage:  python make_worklist_M17.py <Month_RECONCILED_M16.xlsx> <output.xlsx>
Needs:  pip install openpyxl
"""
import sys
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
HI_FILL  = PatternFill("solid", fgColor="FFC7CE")   # high priority (red-ish)
OK_FILL  = PatternFill("solid", fgColor="C6EFCE")   # section headers (green-ish)
TOT_FONT = Font(bold=True)
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))

def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)

def autowidth(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def main(src, dst):
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    ix = {h: i for i, h in enumerate(hdr)}
    rows = [r for r in it if r is not None]

    def g(r, col):
        return r[ix[col]] if col in ix and ix[col] < len(r) else None

    # -- classify ----------------------------------------------------------
    recovery, esi = [], []
    for r in rows:
        if str(g(r, 'ACTION_NEEDED') or '').strip().upper() != 'Y':
            continue
        reason = str(g(r, 'ACTION_REASON') or '').upper()
        if 'ESI-EXEMPT' in reason:
            esi.append(r)
        elif 'PAID ABOVE' in reason or 'IMPLIED' in reason:
            recovery.append(r)

    out = Workbook()

    # -- Tab 2: RECOVERY_REVIEW --------------------------------------------
    rec_cols = ['PRIORITY', 'EMPCODE', 'FULLNAME', 'CLIENTGROUPNAME', 'SITENAME',
                'SITESTATE', 'DESIGNATIONNAME', 'NORMALDAYS', 'FIXEDGROSS',
                'GROSS AMT', 'OT AMOUNT', 'ATTENDANCE ALW', 'EXCESS_TO_REVIEW',
                'ACTION_REASON', 'VERIFIED (Y/N)', 'DECISION (RECOVER/WAIVE/JUSTIFIED)',
                'RECOVERY_MONTH', 'REMARKS']
    wr = out.active
    wr.title = "RECOVERY_REVIEW"
    wr.append(rec_cols)
    style_header(wr, len(rec_cols))
    recovery.sort(key=lambda r: -min(num(g(r, 'EXCESS_SALARY')), num(g(r, 'GROSS AMT'))))
    n_hi = 0
    for r in recovery:
        exc = round(min(num(g(r, 'EXCESS_SALARY')), num(g(r, 'GROSS AMT'))), 0)
        hi = exc > 10000
        n_hi += hi
        wr.append(['HIGH' if hi else 'LOW', g(r, 'EMPCODE'), g(r, 'FULLNAME'),
                   g(r, 'CLIENTGROUPNAME'), g(r, 'SITENAME'), g(r, 'SITESTATE'),
                   g(r, 'DESIGNATIONNAME'), num(g(r, 'NORMALDAYS')),
                   num(g(r, 'FIXEDGROSS')), num(g(r, 'GROSS AMT')),
                   num(g(r, 'OT AMOUNT')), num(g(r, 'ATTENDANCE ALW')), exc,
                   str(g(r, 'ACTION_REASON') or '')[:60], '', '', '', ''])
        if hi:
            for c in range(1, len(rec_cols) + 1):
                wr.cell(row=wr.max_row, column=c).fill = HI_FILL
    autowidth(wr, [8, 11, 24, 26, 24, 14, 18, 7, 10, 10, 9, 9, 12, 40, 9, 16, 12, 24])
    rec_total = sum(min(num(g(r, 'EXCESS_SALARY')), num(g(r, 'GROSS AMT'))) for r in recovery)
    hi_total = sum(min(num(g(r, 'EXCESS_SALARY')), num(g(r, 'GROSS AMT')))
                   for r in recovery if min(num(g(r, 'EXCESS_SALARY')), num(g(r, 'GROSS AMT'))) > 10000)

    # -- Tab 3: ESI_ENROLLMENT ---------------------------------------------
    esi_cols = ['EMPCODE', 'FULLNAME', 'CLIENTGROUPNAME', 'SITENAME', 'SITESTATE',
                'ESIC NO', 'UAN NO', 'GROSS AMT', 'REAL_FULL_MONTH_GROSS',
                'ESI_EE_0.75%', 'ESI_ER_3.25%', 'IC_NO_ALLOTTED',
                'ENROLLED_FROM (month)', 'REMARKS']
    we = out.create_sheet("ESI_ENROLLMENT")
    we.append(esi_cols)
    style_header(we, len(esi_cols))
    esi.sort(key=lambda r: (str(g(r, 'CLIENTGROUPNAME') or ''), -num(g(r, 'GROSS AMT'))))
    esi_gross = 0.0
    for r in esi:
        gr = num(g(r, 'GROSS AMT'))
        esi_gross += gr
        we.append([g(r, 'EMPCODE'), g(r, 'FULLNAME'), g(r, 'CLIENTGROUPNAME'),
                   g(r, 'SITENAME'), g(r, 'SITESTATE'), g(r, 'ESIC NO'),
                   g(r, 'UAN NO'), gr, num(g(r, 'REAL_FULL_MONTH_GROSS')),
                   round(gr * 0.0075, 0), round(gr * 0.0325, 0), '', '', ''])
    autowidth(we, [11, 24, 26, 24, 14, 14, 14, 10, 12, 10, 10, 14, 12, 24])

    # -- Tab 1: SUMMARY ------------------------------------------------------
    s = out.create_sheet("SUMMARY", 0)
    s.append(["MONTHLY WORKLIST — RECOVERY REVIEW & ESI ENROLLMENT (Rule M17)"])
    s.cell(row=1, column=1).font = Font(bold=True, size=13)
    s.append([f"Source: {src.split('/')[-1]}  ·  Excess values are M16 cash-capped"])
    s.append([])
    s.append(["WORKSTREAM", "EMPLOYEES", "AMOUNT (Rs.)", "WHERE"])
    for c in range(1, 5):
        s.cell(row=4, column=c).fill = HDR_FILL
        s.cell(row=4, column=c).font = HDR_FONT
    s.append(["Recovery review — HIGH priority (> Rs.10k each)", n_hi, round(hi_total, 0),
              "Tab RECOVERY_REVIEW (red rows, on top)"])
    s.cell(row=5, column=1).fill = HI_FILL
    s.append(["Recovery review — remaining (small amounts)", len(recovery) - n_hi,
              round(rec_total - hi_total, 0), "Tab RECOVERY_REVIEW"])
    s.append(["ESI enrollment needed (eligible, currently exempt)", len(esi),
              round(esi_gross * 0.04, 0), "Tab ESI_ENROLLMENT (amount = 4%/month exposure)"])
    s.append(["TOTAL", len(recovery) + len(esi), "", ""])
    s.cell(row=8, column=1).font = TOT_FONT
    s.cell(row=8, column=2).font = TOT_FONT
    s.append([])
    # by-client tables
    from collections import defaultdict
    def byclient(rowset, valfn):
        d = defaultdict(lambda: [0, 0.0])
        for r in rowset:
            k = str(g(r, 'CLIENTGROUPNAME') or 'UNKNOWN')[:45]
            d[k][0] += 1
            d[k][1] += valfn(r)
        return sorted(d.items(), key=lambda x: -x[1][1])
    s.append(["RECOVERY BY CLIENT (top 10)", "EMPLOYEES", "EXCESS (Rs.)"])
    s.cell(row=s.max_row, column=1).fill = OK_FILL
    s.cell(row=s.max_row, column=1).font = TOT_FONT
    for k, (n, v) in byclient(recovery, lambda r: min(num(g(r, 'EXCESS_SALARY')), num(g(r, 'GROSS AMT'))))[:10]:
        s.append([k, n, round(v, 0)])
    s.append([])
    s.append(["ESI ENROLLMENT BY CLIENT (top 10)", "EMPLOYEES", "4%/MONTH EXPOSURE (Rs.)"])
    s.cell(row=s.max_row, column=1).fill = OK_FILL
    s.cell(row=s.max_row, column=1).font = TOT_FONT
    for k, (n, v) in byclient(esi, lambda r: num(g(r, 'GROSS AMT')) * 0.04)[:10]:
        s.append([k, n, round(v, 0)])
    autowidth(s, [52, 12, 16, 42])

    out.save(dst)
    print(f"worklist written: {dst}")
    print(f"  RECOVERY_REVIEW : {len(recovery):,} rows (HIGH {n_hi}, Rs.{hi_total:,.0f}) · total Rs.{rec_total:,.0f}")
    print(f"  ESI_ENROLLMENT  : {len(esi):,} rows · 4%/month exposure Rs.{esi_gross*0.04:,.0f}")

if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
