#!/usr/bin/env python3
"""
Rule M16 — Excess-salary projection-artifact patch.

Fixes the two defects found in the April-2026 triage
(docs/pf-salary-reconciliation/april-2026/EXCESS_SALARY_TRIAGE_April2026.md):

  M16a  CAP:      EXCESS_SALARY can never exceed GROSS AMT (you cannot overpay
                  someone more than you actually paid them). The uncapped value
                  is preserved in EXCESS_SALARY_UNCAPPED.
  M16b  SUPPRESS: rows with NORMALDAYS <= 2 flagged "LOW ATTENDANCE" are
                  projection artifacts (full-month projection divides by a tiny
                  day-count) -> ACTION_NEEDED set to 'N', reason annotated
                  [M16_ARTIFACT_LOW_DAYS]. Audit trail kept in RECON_FLAG_M16.

Usage (run locally, needs: pip install openpyxl):
  python patch_excess_artifact_M16.py <input.xlsx> <output.xlsx>

Apply to every future month's reconciled file, or fold the same two rules into
the pipeline that computes EXCESS_SALARY/ACTION_NEEDED.
"""
import sys
import openpyxl
from openpyxl import Workbook

def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0

def main(src, dst):
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    ix = {h: i for i, h in enumerate(hdr)}
    for col in ('EXCESS_SALARY', 'ACTION_NEEDED', 'ACTION_REASON', 'NORMALDAYS', 'GROSS AMT'):
        if col not in ix:
            sys.exit(f"column not found: {col}")

    out = Workbook(write_only=True)
    ows = out.create_sheet("Sheet1")
    ows.append(hdr + ['EXCESS_SALARY_UNCAPPED', 'RECON_FLAG_M16'])

    n = capped = suppressed = 0
    tot_before = tot_after = act_before = act_after = 0.0
    n_act_before = n_act_after = 0
    for r in it:
        if r is None:
            continue
        r = list(r) + [None] * (len(hdr) - len(r))
        n += 1
        excess = num(r[ix['EXCESS_SALARY']])
        gross = num(r[ix['GROSS AMT']])
        days = num(r[ix['NORMALDAYS']])
        action = str(r[ix['ACTION_NEEDED']] or '').strip().upper()
        reason = str(r[ix['ACTION_REASON']] or '')
        flag = []
        uncapped = excess

        if action == 'Y':
            n_act_before += 1
            act_before += excess
        tot_before += excess

        # M16a — cap at gross
        if excess > gross:
            excess = round(gross, 2)
            r[ix['EXCESS_SALARY']] = excess
            flag.append('M16a_CAPPED_AT_GROSS')
            capped += 1
        # M16b — suppress low-day artifact flags
        if action == 'Y' and days <= 2 and 'LOW ATTENDANCE' in reason.upper():
            r[ix['ACTION_NEEDED']] = 'N'
            r[ix['ACTION_REASON']] = reason + ' [M16_ARTIFACT_LOW_DAYS]'
            flag.append('M16b_SUPPRESSED_LOW_DAYS')
            suppressed += 1
            action = 'N'

        if action == 'Y':
            n_act_after += 1
            act_after += excess
        tot_after += excess
        ows.append(r + [uncapped, ';'.join(flag) or None])

    out.save(dst)
    print(f"rows processed          : {n:,}")
    print(f"M16a capped rows        : {capped:,}")
    print(f"M16b suppressed flags   : {suppressed:,}")
    print(f"EXCESS total  before/after : Rs.{tot_before:,.0f} -> Rs.{tot_after:,.0f}")
    print(f"ACTION=Y rows before/after : {n_act_before:,} -> {n_act_after:,}")
    print(f"ACTION=Y excess before/after: Rs.{act_before:,.0f} -> Rs.{act_after:,.0f}")

if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
