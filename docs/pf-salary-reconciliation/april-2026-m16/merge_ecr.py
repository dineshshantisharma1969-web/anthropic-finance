#!/usr/bin/env python3
"""
Merge April-2026 PF (ECR) files -> single per-employee ECR table.

Per vault rules (20-Rules/pf-ecr-case-rules, pf-contribution-basis) and the
pf-salary-reconciliation skill:
  * dedupe EACH file by EMP CODE first (keep first) -- guards against the
    duplicate-row trap that silently doubles the ECR total
  * then concat across files and SUM EE by EMP CODE
    (different files = different PF registrations, so summing is correct)
  * EMPCODE cleaned: strip whitespace, drop any ".0" float suffix

Drop every FORMAT*APRIL*2026*.xlsx into this folder and run.
"""
import glob, os, sys
import openpyxl

FILES = sorted(glob.glob('*.xlsx'))
EXPECTED = ['DELHI', 'STEAGE', 'DMART']


def num(v):
    try:
        return float(str(v).replace(',', ''))
    except Exception:
        return 0.0


def clean_code(v):
    if v is None:
        return ''
    s = str(v).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s.split('.')[0]


def load(path):
    """Return (rows, idx, sheet_total_ee). Header row is auto-detected:
    these files are not consistent (STEAGE has a blank leading row)."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    allrows = list(ws.iter_rows(values_only=True))
    wb.close()
    hrow = None
    for i, r in enumerate(allrows[:20]):
        cells = [str(c).strip().upper() if c is not None else '' for c in r]
        if 'EMP CODE' in cells and 'EE' in cells:
            hrow = i
            break
    if hrow is None:
        raise SystemExit(f'{path}: could not locate header row')
    hdr = [str(h).strip() if h is not None else '' for h in allrows[hrow]]
    idx = {h: i for i, h in enumerate(hdr)}
    allrows = allrows[hrow:]
    for req in ('EMP CODE', 'EE'):
        if req not in idx:
            raise SystemExit(f'{path}: missing column {req!r}. Found: {hdr}')
    data, total_row_ee = [], None
    for r in allrows[1:]:
        code = r[idx['EMP CODE']] if idx['EMP CODE'] < len(r) else None
        if code in (None, ''):
            # Footer rows. The genuine TOTAL row is the first one that carries
            # numbers across several money columns; later footer rows are
            # single-value lines (admin charges, "NET CHALLAN") whose figure can
            # land in the EE column and must NOT be read as the EE total.
            if total_row_ee is None:
                filled = sum(
                    1 for c in ('GROSS', 'Basic Wage', 'PF WAGES', 'EE', 'EPS')
                    if c in idx and idx[c] < len(r) and num(r[idx[c]]) > 0
                )
                if filled >= 3:
                    total_row_ee = num(r[idx['EE']])
            continue
        data.append(r)
    return data, idx, total_row_ee


merged = {}      # empcode -> {'ee':x, 'files':[], 'name':..., 'uan':...}
report = []

for path in FILES:
    up = os.path.basename(path).upper()
    tag = next((e for e in EXPECTED if e in up), os.path.basename(path))
    try:
        data, idx, sheet_ee = load(path)
    except SystemExit as e:
        print(e); continue

    raw_rows = len(data)
    raw_ee = sum(num(r[idx['EE']]) for r in data)

    # --- dedupe within file by EMP CODE (keep first) ---
    seen, ded = set(), []
    for r in data:
        c = clean_code(r[idx['EMP CODE']])
        if c in seen:
            continue
        seen.add(c)
        ded.append(r)
    ded_ee = sum(num(r[idx['EE']]) for r in ded)
    dupes = raw_rows - len(ded)

    tie = ('n/a' if sheet_ee is None
           else ('OK' if abs(raw_ee - sheet_ee) < 1 else f'MISMATCH ({sheet_ee:,.0f})'))
    report.append((tag, raw_rows, dupes, len(ded), raw_ee, ded_ee, tie))

    for r in ded:
        c = clean_code(r[idx['EMP CODE']])
        e = num(r[idx['EE']])
        if c not in merged:
            merged[c] = {'ee': 0.0, 'files': [], 'name': '', 'uan': ''}
        merged[c]['ee'] += e          # sum across files (separate registrations)
        merged[c]['files'].append(tag)
        if not merged[c]['name'] and 'Name' in idx and r[idx['Name']]:
            merged[c]['name'] = str(r[idx['Name']]).strip()
        if not merged[c]['uan'] and 'UAN NO' in idx and r[idx['UAN NO']]:
            merged[c]['uan'] = clean_code(r[idx['UAN NO']])

print('=' * 96)
print('PER-FILE LOAD & DEDUPE')
print('=' * 96)
print(f"{'file':10} {'rows':>7} {'dupes':>7} {'kept':>7} {'raw EE':>16} {'dedup EE':>16}  footer-tie")
for t, rr, dp, kp, re_, de, tie in report:
    print(f'{t:10} {rr:>7,} {dp:>7,} {kp:>7,} {re_:>16,.0f} {de:>16,.0f}  {tie}')

tot_ee = sum(v['ee'] for v in merged.values())
print('-' * 96)
print(f"{'MERGED':10} {'':>7} {'':>7} {len(merged):>7,} {'':>16} {tot_ee:>16,.0f}")
print()
present = [t for t, *_ in report]
missing = [e for e in EXPECTED if e not in present]
if missing:
    print(f'!! INCOMPLETE - missing PF file(s): {", ".join(missing)}')
    print('   The merged total above is NOT the final ECR figure.')
else:
    print('All three PF files present - merged ECR total is complete.')

multi = {k: v for k, v in merged.items() if len(v['files']) > 1}
print(f'\nemployees appearing in >1 PF file (EE summed): {len(multi)}')
for k, v in list(multi.items())[:10]:
    print(f'   {k:>10} {v["name"][:24]:24} {v["files"]}  EE {v["ee"]:,.0f}')

with open('ECR_MERGED_April2026.csv', 'w') as f:
    f.write('EMP_CODE,NAME,UAN_NO,ECR_PF_EE,SOURCE_FILES\n')
    for k in sorted(merged):
        v = merged[k]
        nm = v['name'].replace(',', ' ')
        f.write(f'{k},{nm},{v["uan"]},{v["ee"]:.0f},"{"+".join(v["files"])}"\n')
print(f'\nwrote ECR_MERGED_April2026.csv  ({len(merged):,} employees, EE {tot_ee:,.0f})')
