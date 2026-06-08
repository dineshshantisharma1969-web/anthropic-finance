import docx, copy, os

# ---------- helpers ----------
def set_para(p, text):
    """Set paragraph text, keep first run's formatting."""
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.add_run(text)

def inr(n):
    """Indian comma grouping, integer."""
    n = int(round(n))
    s = str(abs(n))
    if len(s) <= 3:
        out = s
    else:
        last3 = s[-3:]
        rest = s[:-3]
        parts = []
        while len(rest) > 2:
            parts.insert(0, rest[-2:]); rest = rest[:-2]
        if rest: parts.insert(0, rest)
        out = ','.join(parts) + ',' + last3
    return ('-' if n < 0 else '') + out

ONES = ['','One','Two','Three','Four','Five','Six','Seven','Eight','Nine','Ten',
        'Eleven','Twelve','Thirteen','Fourteen','Fifteen','Sixteen','Seventeen','Eighteen','Nineteen']
TENS = ['','','Twenty','Thirty','Forty','Fifty','Sixty','Seventy','Eighty','Ninety']
def two(n):
    if n < 20: return ONES[n]
    return TENS[n//10] + ((' ' + ONES[n%10]) if n%10 else '')
def three(n):
    h = n//100; r = n%100
    out = ''
    if h: out += ONES[h] + ' Hundred'
    if r: out += (' ' if out else '') + two(r)
    return out
def words_in(n):
    n = int(round(n))
    if n == 0: return 'Zero'
    crore = n//10000000; n%=10000000
    lakh = n//100000; n%=100000
    thou = n//1000; n%=1000
    hund = n
    parts = []
    if crore: parts.append(three(crore)+' Crore')
    if lakh: parts.append(three(lakh)+' Lakh')
    if thou: parts.append(three(thou)+' Thousand')
    if hund: parts.append(three(hund))
    return ' '.join(parts)

# ---------- PO records ----------
SUP = dict(hsn='998311', desc='Supervisory Cleaning Services Charges For Mechanized Cleaning at various sites of Ministry of Railway')
HIRE = dict(hsn='997319', desc='Machine Cleaning Hire Charges For Mechanized Cleaning at various sites of Ministry of Railway')

recs = [
 ('01','81','H',1200000,'10-Sep-2023','APR-23 TO JUNE-23'),
 ('02','76','S', 720000,'10-Sep-2023','OCT-23 TO DEC-23'),
 ('03','77','S', 720000,'10-Sep-2023','JAN-24 TO MAR-24'),
 ('04','80','H',1200000,'10-Sep-2023','APR-23 TO JUNE-23'),
 ('05','135-A','H',1680000,'27-Mar-2024','OCT-23 TO DEC-23'),
 ('06','135-B','H',2520000,'27-Mar-2024','JAN-24 TO MAR-24'),
 ('07','135-C','S',1920000,'27-Mar-2024','OCT-23 TO DEC-23'),
 ('08','135-D','S',2160000,'27-Mar-2024','JAN-24 TO MAR-24'),
 ('09','135-E','H',2100000,'27-Mar-2024','OCT-23 TO DEC-23'),
 ('10','135-F','H',2940000,'27-Mar-2024','JAN-24 TO MAR-24'),
 ('11','135-G','S',1680000,'27-Mar-2024','OCT-23 TO DEC-23'),
 ('12','135-H','S',1440000,'27-Mar-2024','JAN-24 TO MAR-24'),
]

os.makedirs('/root/po_out', exist_ok=True)

for seq, inv, typ, taxable, podate, period in recs:
    info = HIRE if typ=='H' else SUP
    igst = round(taxable*0.18)
    gross = taxable + igst
    po_no = f'{seq}/2023-24/UNBS'
    d = docx.Document('/root/blank.docx')

    # T0 C0 : Impressions -> Delhi
    c = d.tables[0].rows[0].cells[0]
    set_para(c.paragraphs[1], 'Branch : DELHI')
    set_para(c.paragraphs[4], '8/7, IST FLOOR, PEELI KOTHI, KIRTI NAGAR INDUSTRIAL AREA')
    set_para(c.paragraphs[5], 'KIRTI NAGAR, WEST DELHI, DELHI - 110015')
    set_para(c.paragraphs[7], 'GST No. : 07AAACI9641L1Z8  |  GST State : Delhi (07)')

    # T0 C1 : PO meta
    c = d.tables[0].rows[0].cells[1]
    set_para(c.paragraphs[2], f'PO No           :  {po_no}')
    set_para(c.paragraphs[3], f'PO Date        :  {podate}')
    set_para(c.paragraphs[4], 'Payment Terms :  AS MUTUALLY AGREED')

    # T1 C0 : Supplier
    c = d.tables[1].rows[0].cells[0]
    set_para(c.paragraphs[0], 'Supplier  :  UDHAS NATH BABA SERVICES PVT LTD')
    set_para(c.paragraphs[2], '                  GT ROAD, ALWARPUR CHOWK, NEAR GYANI MISTRI')
    set_para(c.paragraphs[3], '                  PALWAL, HARYANA - 121102')
    set_para(c.paragraphs[4], '                  GSTIN : 06AABCU5283B2ZJ')
    set_para(c.paragraphs[5], 'GST State :  Haryana (06)')

    # T2 : Ship To (Delhi)
    set_para(d.tables[2].rows[0].cells[0].paragraphs[0],
             'Ship To  :  IMPRESSIONS SERVICES PVT LTD')
    c = d.tables[2].rows[1].cells[0]
    set_para(c.paragraphs[1], '                  8/7, IST FLOOR, PEELI KOTHI, KIRTI NAGAR INDUSTRIAL AREA,')
    set_para(c.paragraphs[2], '                  KIRTI NAGAR, WEST DELHI, DELHI - 110015')

    # T3 : line item row 1
    row = d.tables[3].rows[1].cells
    vals = ['1', f'{info["desc"]} for the period {period}', info['hsn'], '1', 'Nos',
            inr(taxable), '-', '-', '-', '-', '18%', inr(igst), inr(taxable)]
    for ci, v in enumerate(vals):
        set_para(row[ci].paragraphs[0], v)

    # T4 : words + grand total (gross)
    set_para(d.tables[4].rows[0].cells[0].paragraphs[1],
             'Rupees ' + words_in(gross) + ' Only')
    set_para(d.tables[4].rows[0].cells[1].paragraphs[1], f'Rs. {inr(gross)}')

    fname = f'/root/po_out/PO_{seq}_INV-{inv}_{"HIRE" if typ=="H" else "SUPERVISORY"}.docx'
    d.save(fname)
    print(f'{po_no} | {podate} | {"HIRE" if typ=="H" else "SUPV"} | inv {inv} | taxable {inr(taxable)} | igst {inr(igst)} | gross {inr(gross)} | {words_in(gross)}')

print('\nGenerated', len(recs), 'files in /root/po_out')
