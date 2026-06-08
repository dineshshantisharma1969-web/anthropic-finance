import docx, copy, glob, os
from docx.oxml.ns import qn

files = sorted(glob.glob('/root/po_out/PO_*.docx'))
# order by sequence number in filename
files.sort(key=lambda f: int(os.path.basename(f).split('_')[1]))

master = docx.Document(files[0])
mbody = master.element.body

def body_children(doc):
    # all children except the final sectPr
    kids = list(doc.element.body)
    if kids and kids[-1].tag == qn('w:sectPr'):
        kids = kids[:-1]
    return kids

for f in files[1:]:
    master.add_page_break()
    sub = docx.Document(f)
    sectPr = mbody.find(qn('w:sectPr'))
    for el in body_children(sub):
        mbody.insert(list(mbody).index(sectPr), copy.deepcopy(el))

out='/root/UNBS_POs_ALL_12.docx'
master.save(out)
print('combined size', os.path.getsize(out), 'bytes')
# verify count of PO No occurrences
d=docx.Document(out)
n=sum(1 for t in d.tables for r in t.rows for c in r.cells if 'IMP/PO/UNBS' in c.text)
print('PO blocks found:', n, 'tables:', len(d.tables))
